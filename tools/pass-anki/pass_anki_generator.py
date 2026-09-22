#!/usr/bin/env python3
from __future__ import annotations
import argparse, base64, concurrent.futures as cf, hashlib, html, json, os, re, shutil, subprocess, time
import urllib.request
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from pptx import Presentation
from PIL import Image, ImageDraw
import pytesseract

HERE=Path(__file__).resolve().parent
PROMPT=(HERE/"PROMPT.md").read_text(encoding="utf-8")
REVIEW="""Edit these PASS Anki candidates. Keep exactly one memory per card and answers ideally 1-8 words. Split combined answers. Remove exact and near duplicates. Preserve useful reverse cards, Tags, SourceSlide, and exact visible-label answers for image-occlusion cards. Add no facts. Output only: Front | Back | Tags | SourceSlide."""

@dataclass(frozen=True)
class Box: text:str; left:int; top:int; width:int; height:int
@dataclass(frozen=True)
class Slide: number:int; text:str; image:Path|None; width:int; height:int; boxes:tuple[Box,...]
@dataclass(frozen=True)
class Card: front:str; back:str; tags:str; slide:int

def cmd(a): return subprocess.run(a,text=True,capture_output=True,check=True)
def norm(s): return re.sub(r"[^a-z0-9]+"," ",s.casefold().replace("’","'")).strip()
def compact(s): return re.sub(r"\s+"," ",s.strip())

def convert(src,work):
    work.mkdir(parents=True,exist_ok=True)
    if src.suffix.lower()==".pptx":
        dst=work/src.name; shutil.copy2(src,dst); return dst
    if src.suffix.lower()!=".ppt": raise ValueError("Input must be .ppt or .pptx")
    office=shutil.which("soffice") or shutil.which("libreoffice")
    if not office: raise RuntimeError("LibreOffice is required")
    r=cmd([office,"--headless","--convert-to","pptx","--outdir",str(work),str(src)])
    dst=work/(src.stem+".pptx")
    if not dst.exists(): raise RuntimeError(r.stderr or r.stdout)
    return dst

def render(pptx,work):
    office=shutil.which("soffice") or shutil.which("libreoffice"); ppm=shutil.which("pdftoppm")
    if not office or not ppm:return {}
    pdfdir=work/"pdf"; imgdir=work/"images"; pdfdir.mkdir(exist_ok=True); imgdir.mkdir(exist_ok=True)
    cmd([office,"--headless","--convert-to","pdf","--outdir",str(pdfdir),str(pptx)])
    pdf=pdfdir/(pptx.stem+".pdf")
    if not pdf.exists():return {}
    cmd([ppm,"-jpeg","-jpegopt","quality=88","-r","140",str(pdf),str(imgdir/"slide")])
    out={}
    for p in imgdir.glob("slide-*.jpg"):
        m=re.search(r"-(\d+)\.jpg$",p.name)
        if m:out[int(m.group(1))]=p
    return out

def extract(pptx,images):
    prs=Presentation(str(pptx)); out=[]
    for n,sl in enumerate(prs.slides,1):
        texts=[]; boxes=[]
        for sh in sl.shapes:
            t=compact(getattr(sh,"text",""))
            if t:texts.append(t); boxes.append(Box(t,int(sh.left),int(sh.top),int(sh.width),int(sh.height)))
        out.append(Slide(n,"\n".join(texts),images.get(n),int(prs.slide_width),int(prs.slide_height),tuple(boxes)))
    return out

def data_url(p):return "data:image/jpeg;base64,"+base64.b64encode(p.read_bytes()).decode()
def api(messages,key,base,model,temp=.1,max_tokens=12000):
    url=base.rstrip("/")+"/chat/completions"; payload=json.dumps({"model":model,"messages":messages,"temperature":temp,"max_tokens":max_tokens}).encode()
    for attempt in range(5):
        req=urllib.request.Request(url,data=payload,method="POST",headers={"Authorization":f"Bearer {key}","Content-Type":"application/json"})
        try:
            with urllib.request.urlopen(req,timeout=300) as r:obj=json.loads(r.read().decode())
            return obj["choices"][0]["message"]["content"]
        except Exception:
            if attempt==4:raise
            time.sleep(2**attempt)

def messages(batch,with_images):
    content=[{"type":"text","text":PROMPT+"\nCreate exhaustive cards from these slides."}]
    for s in batch:
        labels="\n".join("- "+b.text for b in s.boxes if len(b.text)<=120) or "[none]"
        content.append({"type":"text","text":f"\n--- SLIDE {s.number} ---\nTEXT:\n{s.text or '[none]'}\nTEXT SHAPES FOR OCCLUSION:\n{labels}"})
        if with_images and s.image:content.append({"type":"image_url","image_url":{"url":data_url(s.image),"detail":"high"}})
    return [{"role":"user","content":content}]

def parse(text):
    cards=[]; rejects=[]
    for raw in text.replace("```text","").replace("```","").splitlines():
        line=raw.strip()
        if not line:continue
        if line.count("|")!=3:rejects.append(line);continue
        f,b,t,s=(compact(x) for x in line.split("|",3)); m=re.search(r"\d+",s); tags=" ".join(dict.fromkeys(re.findall(r"[A-Za-z0-9_:-]+",t)))
        if not f or not b or not tags or not m or len(f)>500 or len(b)>180:rejects.append(line);continue
        cards.append(Card(f,b,tags,int(m.group())))
    return cards,rejects

def dedupe(cards):
    seen=set();out=[]
    for c in cards:
        k=(norm(c.front),norm(c.back))
        if k not in seen:seen.add(k);out.append(c)
    return out

def review(cards,key,base,model):
    out=[];bad=[]
    for i in range(0,len(cards),150):
        payload="\n".join(f"{c.front} | {c.back} | {c.tags} | {c.slide}" for c in cards[i:i+150])
        good,reject=parse(api([{"role":"user","content":REVIEW+"\n\n"+payload}],key,base,model,0,14000));out+=good;bad+=reject
    return dedupe(out),bad

def shape_box(s,answer):
    target=norm(answer);exact=[b for b in s.boxes if norm(b.text)==target]
    if exact:return min(exact,key=lambda b:b.width*b.height)
    near=[b for b in s.boxes if target in norm(b.text) and len(norm(b.text))<=len(target)+30]
    return min(near,key=lambda b:b.width*b.height) if near else None

@lru_cache(maxsize=256)
def ocr(path):
    d=pytesseract.image_to_data(Image.open(path).convert("RGB"),output_type=pytesseract.Output.DICT,config="--psm 11")
    return tuple((t.strip(),int(d["left"][i]),int(d["top"][i]),int(d["width"][i]),int(d["height"][i])) for i,t in enumerate(d["text"]) if t.strip() and float(d["conf"][i])>=20)
def ocr_box(path,answer):
    target=re.sub(r"[^a-z0-9]+","",answer.casefold());best=None
    if not target:return None
    try:words=ocr(str(path))
    except Exception:return None
    for i in range(len(words)):
        for length in range(1,min(7,len(words)-i)+1):
            g=words[i:i+length];candidate=re.sub(r"[^a-z0-9]+","","".join(x[0] for x in g).casefold())
            if candidate==target or (len(target)>=4 and target in candidate and len(candidate)<=len(target)+6):
                x1=min(x[1] for x in g);y1=min(x[2] for x in g);x2=max(x[1]+x[3] for x in g);y2=max(x[2]+x[4] for x in g);area=(x2-x1)*(y2-y1)
                if best is None or area<best[0]:best=(area,x1,y1,x2,y2)
    return best[1:] if best else None

def mask(s,target,media,c):
    im=Image.open(s.image).convert("RGB")
    if isinstance(target,Box):
        sx,sy=im.width/s.width,im.height/s.height;x1,y1=int(target.left*sx),int(target.top*sy);x2,y2=int((target.left+target.width)*sx),int((target.top+target.height)*sy)
    else:x1,y1,x2,y2=target
    pad=max(3,int(im.width*.003));x1=max(0,x1-pad);y1=max(0,y1-pad);x2=min(im.width,x2+pad);y2=min(im.height,y2+pad)
    ImageDraw.Draw(im).rectangle((x1,y1,x2,y2),fill=(255,241,140),outline=(40,40,40),width=max(2,pad//2))
    name=f"occ_{s.number:03d}_{hashlib.sha1((c.front+c.back).encode()).hexdigest()[:10]}.jpg";im.save(media/name,quality=92);return name

def write(output,cards,rejects,slides):
    output.mkdir(parents=True,exist_ok=True);media=output/"media";media.mkdir(exist_ok=True);sm={s.number:s for s in slides};copied={};rows=[]
    for c in cards:
        front=c.front;tags=set(c.tags.split());s=sm.get(c.slide)
        if "image" in tags:
            if not s or not s.image:rejects.append(f"Missing image | {c.front} | {c.slide}");continue
            name=""
            if "image-occlusion" in tags:
                target=shape_box(s,c.back) or ocr_box(s.image,c.back)
                if target:name=mask(s,target,media,c)
                else:rejects.append(f"Occlusion target not found; full image used | {c.front} | {c.back} | {c.slide}")
            if not name:
                name=copied.get(s.number,"")
                if not name:name=f"slide_{s.number:03d}.jpg";shutil.copy2(s.image,media/name);copied[s.number]=name
            front=f'<img src="{name}"><br>{html.escape(c.front)}'
        rows.append((front,c.back,c.tags+f" slide::{c.slide:03d}"))
    (output/"cards.tsv").write_text("\n".join("\t".join(r) for r in rows)+"\n",encoding="utf-8")
    (output/"cards.txt").write_text("\n".join(" | ".join(r) for r in rows)+"\n",encoding="utf-8")
    (output/"review_rejects.txt").write_text("\n".join(rejects)+( "\n" if rejects else ""),encoding="utf-8")
    (output/"extracted_slides.txt").write_text("\n".join(f"===== SLIDE {s.number} =====\n{s.text}\n" for s in slides),encoding="utf-8")

def main():
    p=argparse.ArgumentParser();p.add_argument("powerpoint",type=Path);p.add_argument("--output",type=Path,default=Path("anki_output"));p.add_argument("--model",default=os.getenv("OPENAI_MODEL","gpt-4.1-mini"));p.add_argument("--base-url",default=os.getenv("OPENAI_BASE_URL","https://api.openai.com/v1"));p.add_argument("--batch-size",type=int,default=3);p.add_argument("--workers",type=int,default=2);p.add_argument("--text-only",action="store_true");p.add_argument("--skip-review",action="store_true");a=p.parse_args()
    src=a.powerpoint.resolve();work=a.output.resolve()/"_work";pptx=convert(src,work);images={} if a.text_only else render(pptx,work);slides=extract(pptx,images);key=os.getenv("OPENAI_API_KEY")
    if not key:raise SystemExit("OPENAI_API_KEY is required")
    batches=[slides[i:i+a.batch_size] for i in range(0,len(slides),a.batch_size)]
    def generate(x):
        i,b=x;good,bad=parse(api(messages(b,not a.text_only),key,a.base_url,a.model));print(f"Batch {i+1}/{len(batches)}: {len(good)} cards");return i,good,bad
    with cf.ThreadPoolExecutor(max_workers=a.workers) as pool:results=list(pool.map(generate,enumerate(batches)))
    cards=[];rejects=[]
    for _,good,bad in sorted(results):cards+=good;rejects+=bad
    cards=dedupe(cards)
    if cards and not a.skip_review:cards,bad=review(cards,key,a.base_url,a.model);rejects+=bad
    write(a.output.resolve(),cards,rejects,slides);print(f"Wrote {len(cards)} cards to {a.output}")
if __name__=="__main__":main()
